"""Generate Python oracle and XLSX fixtures for browser-parity.mjs."""
import json
import sys
from dataclasses import asdict
from datetime import date, time
from types import SimpleNamespace
from pathlib import Path
from openpyxl import Workbook, load_workbook
from miniplan.readers.xlsx import read_xlsx
from miniplan.pipeline import _load_calendars
from miniplan.classify import ServiceClassifier
from miniplan.calendar import HolidayCalendar
from miniplan.schedule import build_schedule
from miniplan.writer import write_schedule

reference, output = map(Path, sys.argv[1:])
workbook = Workbook()
workbook.active.title = 'Hinweise'
workbook.active['A1'] = 'Ignore this sheet'
sheet = workbook.create_sheet('Plan')
workbook.active = 1
sheet['A1'] = 'Donnerstag, 2. Juli 2026'
sheet.append(['Hochfest'])
for row, start, name in [(4, time(18, 30), 'Eucharistiefeier'), (7, 'morgen', 'Tauffeier')]:
    sheet.cell(row, 2, start)
    sheet.cell(row, 4, name)
    sheet.cell(row, 5, 'Kirche St. Georg')
sheet.cell(8, 1, 'Sonntag, 5. Juli 2026')
sheet.cell(9, 2, '10:00')
sheet.cell(9, 4, 'Wortgottesdienst')
sheet.cell(9, 5, 'Kirche St. Georg')
workbook.save(output / 'input.xlsx')
calendar_path = output / 'custom.ics'
calendar_path.write_text('BEGIN:VCALENDAR\nVERSION:2.0\nBEGIN:VEVENT\nDTSTART;VALUE=DATE:20260702\nDTEND;VALUE=DATE:20260703\nEND:VEVENT\nEND:VCALENDAR\n')
rules = json.loads((reference / 'src/miniplan/config/services.json').read_text())
rules['Eucharistiefeier']['Bezeichnung'] = 'Eigene Messe'
(output / 'rules.json').write_text(json.dumps(rules))

def result(read, classifier):
    entries = build_schedule(read.services, classifier)
    path = write_schedule(entries, output / 'python')
    sheet = load_workbook(path).active
    return {
        'services': [{'date': str(s.date), 'start': s.start.strftime('%H:%M'), 'end': s.end.strftime('%H:%M') if s.end else '', 'name': s.name, 'location': s.location, 'dayInfo': s.day_info} for s in read.services],
        'diagnostics': [{'message': d.message, 'row': d.page_or_row} for d in read.diagnostics],
        'plan': [{**asdict(e), 'date': str(e.date), 'start': e.start.strftime('%H:%M') if e.start else None} for e in entries],
        'rows': list(sheet.iter_rows(values_only=True)),
    }

xlsx = read_xlsx(output / 'input.xlsx')
oracle = {'xlsx': result(xlsx, ServiceClassifier.default(_load_calendars(None))),
          'custom': result(xlsx, ServiceClassifier(rules, HolidayCalendar.from_ics_paths([calendar_path]))),
          'none': result(xlsx, ServiceClassifier(rules, HolidayCalendar()))}
# PDF semantics are asserted independently in JavaScript. Use those corrected
# source services only to retain Python scheduling/writer parity.
pdf_path = output / 'pdf-services.json'
if pdf_path.exists():
    services = [SimpleNamespace(date=date.fromisoformat(s['date']),
        start=time.fromisoformat(s['start']), end=time.fromisoformat(s['end']) if s['end'] else None,
        name=s['name'], location=s['location'], day_info=s['dayInfo'])
        for s in json.loads(pdf_path.read_text())]
    oracle['pdf'] = result(SimpleNamespace(services=services, diagnostics=[]),
        ServiceClassifier.default(_load_calendars(None)))
(output / 'oracle.json').write_text(json.dumps(oracle))
